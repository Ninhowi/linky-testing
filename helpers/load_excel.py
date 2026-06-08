"""Load pytest test data from Excel workbooks (row 1 = headers, row 2+ = data)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

TestRow = dict[str, Any]
SheetRows = list[TestRow]
WorkbookData = dict[str, SheetRows]


def load_excel(path: str | Path, *, data_only: bool = True) -> WorkbookData:
    """Load every sheet in an Excel file as a list of row dicts.

    Format per sheet:
      - Row 1: column names (field headers)
      - Row 2 onward: test data until the last non-empty row

    Returns:
        ``{sheet_name: [{"field": value, ...}, ...], ...}``
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Excel file not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=data_only)
    try:
        return {name: load_sheet(workbook[name]) for name in workbook.sheetnames}
    finally:
        workbook.close()


def load_sheet(worksheet: Worksheet | str | Path, sheet: str | None = None) -> SheetRows:
    """Load one sheet as row dicts.

    Pass either a :class:`Worksheet`, or a file path plus optional ``sheet`` name
    (defaults to the active sheet).
    """
    if isinstance(worksheet, (str, Path)):
        path = Path(worksheet)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            name = sheet or workbook.active.title
            if name not in workbook.sheetnames:
                raise ValueError(
                    f"Sheet {name!r} not in {path.name}. "
                    f"Available: {', '.join(workbook.sheetnames)}"
                )
            return _rows_from_worksheet(workbook[name])
        finally:
            workbook.close()

    return _rows_from_worksheet(worksheet)


def _rows_from_worksheet(ws: Worksheet) -> SheetRows:
    row_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(row_iter)
    except StopIteration:
        return []

    headers = _parse_headers(header_row)
    rows: SheetRows = []
    for row in row_iter:
        if _is_empty_row(row):
            continue
        rows.append(_row_to_dict(headers, row))
    return rows


def _parse_headers(header_row: tuple[Any, ...]) -> list[tuple[int, str]]:
    if not header_row or all(_is_blank(cell) for cell in header_row):
        raise ValueError("Missing header row (row 1 must contain column names)")

    headers: list[tuple[int, str]] = []
    for index, cell in enumerate(header_row):
        if _is_blank(cell):
            continue
        headers.append((index, str(cell).strip()))

    if not headers:
        raise ValueError("Missing header row (row 1 must contain column names)")

    seen: set[str] = set()
    duplicates: list[str] = []
    for _, name in headers:
        if name in seen:
            duplicates.append(name)
        seen.add(name)
    if duplicates:
        raise ValueError(f"Duplicate column names in header row: {duplicates}")
    return headers


def _row_to_dict(headers: list[tuple[int, str]], row: tuple[Any, ...]) -> TestRow:
    return {
        name: row[index] if index < len(row) else None
        for index, name in headers
    }


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return not row or all(_is_blank(cell) for cell in row)


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False
